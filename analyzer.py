from __future__ import annotations

import difflib
import json
import logging
import os
import queue
import re
import shutil
import sys
import tempfile
import threading
import time
import traceback
import zipfile
from concurrent.futures import ThreadPoolExecutor
from contextlib import suppress
from dataclasses import asdict
from pathlib import Path
from tkinter import DoubleVar, StringVar, Tk, filedialog, messagebox
from tkinter import ttk

from app import (
    APP_DIR,
    ChangeGroup,
    OperationCancelled,
    WordEngine,
    close_operation_logger,
    create_comment_free_copy,
    create_operation_logger,
    detect_metadata_from_name,
)


VERSION = "변분기 Alpha 0.4"
PACKAGE_EXTENSION = ".bdsg"
PACKAGE_FORMAT = "byeondaesaenggi-analysis-v1"


def _safe_extract(archive: zipfile.ZipFile, destination: Path) -> None:
    root = destination.resolve()
    for member in archive.infolist():
        target = (destination / member.filename).resolve()
        if root != target and root not in target.parents:
            raise ValueError("분석 패키지 내부 경로가 올바르지 않습니다.")
    archive.extractall(destination)


def _range_text(document, group_index: int, engine: WordEngine) -> str:
    start_name = engine._bookmark_name(group_index, "start")
    end_name = engine._bookmark_name(group_index, "end")
    if not document.Bookmarks.Exists(start_name) or not document.Bookmarks.Exists(end_name):
        raise RuntimeError(f"분석 범위 북마크가 없습니다: {group_index}")
    start = int(document.Bookmarks(start_name).Range.Start)
    end = int(document.Bookmarks(end_name).Range.Start)
    value = str(document.Range(min(start, end), max(start, end)).Text or "")
    # Word table cells end with CR+BEL. Preserve rows and cells in plain text so
    # the preview remains usable without nesting a table in another Word table.
    value = value.replace("\r\x07", "\n").replace("\x07", "\t")
    value = value.replace("\r", "\n").replace("\v", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _with_omission_markers(text: str, group: ChangeGroup) -> str:
    before = "<앞쪽 표 행 생략>" if group.context_type == "table" else "<전략>"
    after = "<뒤쪽 표 행 생략>" if group.context_type == "table" else "<후략>"
    parts = []
    if group.omit_before:
        parts.append(before)
    if text:
        parts.append(text)
    if group.omit_after:
        parts.append(after)
    return "\n".join(parts)


def _bookmark_range(document, group_index: int, engine: WordEngine):
    start_name = engine._bookmark_name(group_index, "start")
    end_name = engine._bookmark_name(group_index, "end")
    if not document.Bookmarks.Exists(start_name) or not document.Bookmarks.Exists(end_name):
        raise RuntimeError(f"분석 범위 북마크가 없습니다: {group_index}")
    start = int(document.Bookmarks(start_name).Range.Start)
    end = int(document.Bookmarks(end_name).Range.Start)
    return document.Range(min(start, end), max(start, end))


def _capture_range_png(document, group_index: int, engine: WordEngine, output: Path):
    from word_capture import capture_range
    return capture_range(_bookmark_range(document, group_index, engine), engine, output)


def excel_text(value):
    if not isinstance(value, str):
        return value
    # Only Excel copies are cleaned/truncated. Full source stays in analysis.json.
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", value)
    return value[:32767]


def _write_preview_xlsx(path: Path, items: list[dict], document_name: str, warnings=None) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as ExcelImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    original_items = items
    items = [{key: excel_text(value) for key, value in item.items()} for item in items]
    document_name = excel_text(document_name)
    # Keep one row per reviewable change range. The page value is repeated when
    # a page has several ranges, which avoids stacking multiple images in one cell.
    page_items = list(items)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "페이지별 변경 미리보기"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:F1")
    sheet["A1"] = excel_text(f"변경 내용 미리보기 · {document_name}")
    sheet["A1"].font = Font(name="맑은 고딕", size=14, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    headers = ["No.", "Section", "page", "변경 유형", "변경 전", "변경 후"]
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(3, column, value)
        cell.font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_number, item in enumerate(page_items, start=4):
        before_image = item.get("before_image")
        after_image = item.get("after_image")
        values = [
            item["index"], item["section"], item["page"],
            "표" if item.get("context_type") == "table" else "문단",
            "" if before_image else item["before"],
            "" if after_image else item["after"],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            cell.font = Font(name="맑은 고딕", size=9)
            if column == 5:
                cell.fill = PatternFill("solid", fgColor="FCE8E6")
            elif column == 6:
                cell.fill = PatternFill("solid", fgColor="E8F1FB")
            cell.alignment = Alignment(
                horizontal="center" if column <= 4 else "left",
                vertical="top",
                wrap_text=True,
            )
        image_heights = []
        for column_index, column_letter, image_path in (
            (4, "E", before_image), (5, "F", after_image)
        ):
            if not image_path or not Path(image_path).exists():
                continue
            picture = ExcelImage(str(image_path))
            scale = min(1.0, 520 / max(1, picture.width), 520 / max(1, picture.height))
            picture.width = max(1, round(picture.width * scale))
            picture.height = max(1, round(picture.height * scale))
            picture.anchor = TwoCellAnchor(
                editAs="twoCell",
                _from=AnchorMarker(col=column_index, row=row_number - 1),
                to=AnchorMarker(col=column_index, row=row_number - 1, colOff=round(picture.width*9525), rowOff=round(picture.height*9525)),
            )
            sheet.add_image(picture)
            image_heights.append(picture.height)
        text_lines = max(item["before"].count("\n") + 1, item["after"].count("\n") + 1)
        text_height = min(180, max(36, text_lines * 14))
        picture_height = max(image_heights, default=0) * 0.75 + (12 if image_heights else 0)
        sheet.row_dimensions[row_number].height = min(405, max(text_height, picture_height))
        # Searchable source text and summary are retained in hidden columns.
        sheet.cell(row_number, 7, item["before"])
        sheet.cell(row_number, 8, item["after"])
        sheet.cell(row_number, 9, item.get("summary", ""))
    widths = [7, 30, 24, 11, 76, 76, 2, 2, 2]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet["G3"] = "변경 전 원문"
    sheet["H3"] = "변경 후 원문"
    sheet["I3"] = "변경 요약"
    for column_letter in ("G", "H", "I"):
        sheet.column_dimensions[column_letter].hidden = True
    sheet.freeze_panes = "E4"
    sheet.auto_filter.ref = f"A3:F{max(3, len(page_items) + 3)}"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    detail = workbook.create_sheet("변경항목별 상세")
    detail.append(["번호", "Section", "page", "형식", "미리보기 방식", "변경 전", "변경 후", "비고"])
    for item in items:
        detail.append([
            item["index"], item["section"], item["page"],
            "표" if item.get("context_type") == "table" else "문단",
            item.get("image_status", "텍스트"),
            item["before"], item["after"],
            ("전후 대응/서식 영향 확인 필요. " if item.get("needs_review") else "") + item.get("image_error", ""),
        ])
    for cell in detail[1]:
        cell.font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in detail.iter_rows(min_row=2):
        for column, cell in enumerate(row, start=1):
            cell.font = Font(name="맑은 고딕", size=9)
            if column == 6:
                cell.fill = PatternFill("solid", fgColor="FCE8E6")
            elif column == 7:
                cell.fill = PatternFill("solid", fgColor="E8F1FB")
            cell.alignment = Alignment(
                horizontal="center" if column <= 5 else "left",
                vertical="top", wrap_text=True,
            )
    for index, width in enumerate([7, 32, 11, 10, 14, 60, 60, 36], start=1):
        detail.column_dimensions[get_column_letter(index)].width = width
    detail.freeze_panes = "F2"
    detail.auto_filter.ref = f"A1:H{max(1, len(items) + 1)}"
    detail.sheet_view.showGridLines = False
    info = workbook.create_sheet("확인사항")
    info.append(["항목", "내용"])
    for warning in warnings or []:
        info.append(["확인 필요", excel_text(warning)])
    info.append(["비교 단위", "본문 문단/최상위 표 블록. 같은 페이지에 여러 행이 있을 수 있습니다."])
    info.append(["이미지", "여러 페이지 블록은 잘림 방지를 위해 텍스트로 대체. 원본 Word를 확인하십시오."])
    info.append(["텍스트 보존", "Excel용 제어문자 정제 및 셀당 32767자 제한. 전체 원문은 analysis.json에 보존."])
    for original in original_items:
        changed = [key for key, value in original.items() if isinstance(value, str) and excel_text(value) != value]
        if changed:
            info.append([f"항목 {original.get('index', '?')}", "Excel 정제/잘림: " + ", ".join(changed)])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 100
    for ws in workbook:
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Document text is never an Excel formula.
    workbook.save(path)
    workbook.close()


def create_analysis_package(before_path, after_path, tracked_path, package_path, progress, cancel_event, logger):
    from pair_runner import create_package
    return create_package(before_path, after_path, tracked_path, package_path, progress, cancel_event, logger,
                          preview_writer=_write_preview_xlsx, capture=_capture_range_png,
                          bookmark_range=_bookmark_range, version=VERSION, package_format=PACKAGE_FORMAT)


def read_package_manifest(package_path: Path) -> dict:
    with zipfile.ZipFile(package_path) as archive:
        if archive.testzip() is not None or "analysis.json" not in archive.namelist():
            raise ValueError("올바른 변대생기 분석 패키지가 아닙니다.")
        manifest = json.loads(archive.read("analysis.json").decode("utf-8"))
    if manifest.get("format") != PACKAGE_FORMAT:
        raise ValueError("지원하지 않는 분석 패키지 형식입니다.")
    return manifest


class SimpleTool:
    def __init__(self, root: Tk) -> None:
        self.root = root
        self.mode = "analyze"
        self.events: queue.Queue = queue.Queue()
        self.cancel_event = threading.Event()
        self.busy = False
        self.input_var = StringVar()
        self.before_var = StringVar()
        self.tracked_var = StringVar()
        self.output_var = StringVar()
        self.status_var = StringVar(value="입력 파일을 선택하십시오.")
        self.stage_var = StringVar(value="대기 중")
        self.detail_var = StringVar(value="진행할 작업이 없습니다.")
        self.elapsed_var = StringVar(value="경과 시간 00:00")
        self.progress_var = DoubleVar(value=0)
        self.last_log_path: Path | None = None
        self.task_started_at = 0.0
        self.last_progress_at = 0.0
        title = "변분기 - Alpha 0.4"
        root.title(title)
        root.geometry("900x540")
        root.minsize(800, 500)
        frame = ttk.Frame(root, padding=18)
        frame.pack(fill="both", expand=True)
        description = "변경 전·후 DOCX는 필수, 변경이력 DOCX는 선택입니다. 전후 원문을 기준으로 분석하고 Word 이미지를 Excel에 배치합니다."
        ttk.Label(frame, text=description, wraplength=800).pack(fill="x", pady=(0, 18))
        self._file_row(frame, "변경 전 *", self.before_var, lambda: self.choose_docx(self.before_var))
        self._file_row(frame, "변경 후 *", self.input_var, lambda: self.choose_docx(self.input_var))
        self._file_row(frame, "이력 선택", self.tracked_var, lambda: self.choose_docx(self.tracked_var))
        self._file_row(frame, "출력 *", self.output_var, self.choose_output)
        self.run_button = ttk.Button(
            frame,
            text="변경 분석 및 미리보기 생성",
            command=self.run,
        )
        self.run_button.pack(pady=(12, 4))
        self.cancel_button = ttk.Button(frame, text="취소", command=self.cancel, state="disabled")
        self.cancel_button.pack(pady=(0, 8))
        root.protocol("WM_DELETE_WINDOW", self.on_close)
        progress_frame = ttk.LabelFrame(frame, text="진행 상황", padding=10)
        progress_frame.pack(fill="x")
        progress_top = ttk.Frame(progress_frame)
        progress_top.pack(fill="x")
        ttk.Label(progress_top, textvariable=self.stage_var).pack(side="left")
        ttk.Label(progress_top, textvariable=self.elapsed_var).pack(side="right")
        ttk.Progressbar(progress_frame, variable=self.progress_var, maximum=100).pack(fill="x", pady=7)
        progress_bottom = ttk.Frame(progress_frame)
        progress_bottom.pack(fill="x")
        ttk.Label(progress_bottom, textvariable=self.detail_var, wraplength=650).pack(side="left", fill="x", expand=True)
        self.log_button = ttk.Button(progress_bottom, text="로그 열기", command=self.open_log, state="disabled")
        self.log_button.pack(side="right", padx=(8, 0))
        ttk.Label(frame, textvariable=self.status_var, wraplength=800).pack(fill="x", pady=(12, 0))

    @staticmethod
    def _file_row(parent, label, variable, command) -> None:
        row = ttk.Frame(parent)
        row.pack(fill="x", pady=5)
        ttk.Label(row, text=label, width=7).pack(side="left")
        ttk.Entry(row, textvariable=variable).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(row, text="찾아보기", command=command).pack(side="right")

    def choose_docx(self, variable):
        name = filedialog.askopenfilename(filetypes=[("Word 문서", "*.docx")])
        if name:
            variable.set(name)
            if variable is self.input_var and not self.output_var.get():
                self.output_var.set(str(Path(name).with_name(Path(name).stem + "_분석.bdsg")))

    def choose_output(self):
        name = filedialog.asksaveasfilename(defaultextension=".bdsg", filetypes=[("변분기 분석 패키지", "*.bdsg")])
        if name:
            self.output_var.set(name)

    def cancel(self):
        self.cancel_event.set()
        self.status_var.set("취소 요청됨: 현재 Word 호출이 끝난 후 중단합니다.")
        self.cancel_button.configure(state="disabled")

    def on_close(self):
        if self.busy:
            self.cancel()
            messagebox.showinfo(self.root.title(), "취소를 요청했습니다. Word 정리가 끝난 후 창을 닫아 주세요.")
        else:
            self.root.destroy()

    def run(self) -> None:
        from pair_analysis import validate_paths
        if not self.before_var.get().strip() or not self.input_var.get().strip() or not self.output_var.get().strip():
            messagebox.showwarning(self.root.title(), "변경 전·후 및 출력 파일을 지정하십시오.")
            return
        before = Path(self.before_var.get().strip())
        source, output = Path(self.input_var.get().strip()), Path(self.output_var.get().strip())
        tracked = Path(self.tracked_var.get().strip()) if self.tracked_var.get().strip() else None
        try:
            outputs = validate_paths(before, source, tracked, output)
        except ValueError as exc:
            messagebox.showwarning(self.root.title(), str(exc))
            return
        if any(p.exists() for p in outputs):
            if not messagebox.askyesno(self.root.title(), "동일 이름의 결과물이 있습니다. 덮어쓸까요?"):
                return
        self.cancel_event.clear()
        self.cancel_button.configure(state="normal")
        self.busy = True
        self.run_button.configure(state="disabled")
        self.progress_var.set(0)
        self.stage_var.set("작업 준비")
        self.detail_var.set(source.name)
        self.task_started_at = time.monotonic()
        self.last_progress_at = self.task_started_at
        self.elapsed_var.set("경과 시간 00:00")
        self.status_var.set("작업을 시작했습니다.")
        logger, log_path = create_operation_logger("변분기_분석")
        self.last_log_path = log_path
        self.log_button.configure(state="normal")

        def record_progress(event: dict) -> None:
            if event.get("kind") == "progress":
                logger.info(
                    "%s %.1f%% (%s/%s) %s",
                    event.get("stage", "처리 중"),
                    float(event.get("percent", 0)),
                    event.get("current", "-"),
                    event.get("total", "-"),
                    event.get("detail", ""),
                )
            self.events.put(event)

        def worker() -> None:
            try:
                logger.info("작업 시작 · 전: %s · 후: %s · 이력: %s · 출력: %s", before, source, tracked, output)
                result = create_analysis_package(before, source, tracked, output, record_progress, self.cancel_event, logger)
                logger.info("작업 완료 · 검토 블록 %s개", result["count"])
                self.events.put({"kind": "done", **result, "output": output})
            except OperationCancelled:
                logger.info("사용자 취소")
                self.events.put({"kind": "cancelled"})
            except Exception as exc:
                logger.exception("작업 오류")
                self.events.put({"kind": "error", "error_type": type(exc).__name__, "message": str(exc)[:240], "log": log_path})
            finally:
                close_operation_logger(logger)

        threading.Thread(target=worker, daemon=True).start()
        self.root.after(120, self.poll)

    def poll(self) -> None:
        terminal = None
        while True:
            try:
                event = self.events.get_nowait()
            except queue.Empty:
                break
            if event.get("kind") == "progress":
                self.last_progress_at = time.monotonic()
                self.progress_var.set(max(self.progress_var.get(), float(event.get("percent", 0))))
                detail = str(event.get("detail", ""))
                current, total = event.get("current"), event.get("total")
                if current is not None and total:
                    detail = f"{current}/{total} · {detail}"
                self.stage_var.set(str(event.get("stage", "처리 중")))
                self.detail_var.set(detail)
            elif event.get("kind") in {"done", "error", "cancelled"}:
                terminal = event
        if terminal:
            self.busy = False
            self.cancel_button.configure(state="disabled")
            self.run_button.configure(state="normal")
            if terminal["kind"] == "done":
                self.progress_var.set(100)
                self.stage_var.set("분석 완료")
                self.detail_var.set(f"변경 항목 {terminal['count']}개 · 분석 패키지 저장 완료")
                output = Path(terminal["output"])
                preview_output = output.with_name(output.stem + "_변경내용_미리보기.xlsx")
                self.status_var.set(f"완료: {output}")
                messagebox.showinfo(
                    self.root.title(),
                    "작업이 완료되었습니다.\n\n"
                    f"변경 항목: {terminal['count']}개\n"
                    f"미리보기: {preview_output}\n"
                    f"분석 패키지: {output}\n"
                    f"확인사항: {len(terminal.get('warnings', []))}개 (Excel 확인사항 시트 참조)",
                )
                with suppress(OSError):
                    os.startfile(preview_output if preview_output.exists() else output)
            elif terminal["kind"] == "cancelled":
                self.stage_var.set("취소됨")
                self.status_var.set("작업을 취소했습니다. 원본은 변경되지 않았습니다.")
            else:
                failed_stage = self.stage_var.get()
                self.stage_var.set("작업 오류")
                self.detail_var.set("로그 열기에서 마지막 완료 단계와 오류 상세를 확인하십시오.")
                self.status_var.set("작업 중 오류가 발생했습니다.")
                messagebox.showerror(self.root.title(), f"오류 유형: {terminal['error_type']}\n발생 단계: {failed_stage}\n{terminal['message']}\n\n로그: {terminal['log']}")
            return
        elapsed = max(0, int(time.monotonic() - self.task_started_at))
        minutes, seconds = divmod(elapsed, 60)
        self.elapsed_var.set(f"경과 시간 {minutes:02d}:{seconds:02d}")
        silent = time.monotonic() - self.last_progress_at
        if silent >= 120:
            self.status_var.set(
                f"마지막 진행 기록 후 {int(silent)}초가 지났습니다. 로그 열기에서 현재 단계를 확인하십시오."
            )
        self.root.after(120, self.poll)

    def open_log(self) -> None:
        if self.last_log_path and self.last_log_path.exists():
            with suppress(OSError):
                os.startfile(self.last_log_path)


def main() -> int:
    if sys.platform != "win32":
        print("변분기는 Microsoft Word가 설치된 Windows에서 실행하십시오.")
        return 2
    root = Tk()
    SimpleTool(root)
    root.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
