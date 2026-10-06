"""Manual side matching must preserve every source block and saved review choice."""
import copy
import hashlib
import io
import json
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.oxml.ns import qn
from PIL import Image
import generator as g


class MatchingReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.docs = {}
        self.images = {}
        texts = {
            'before': ['알파 기준 10', '베타 기준 20', '감마 기준 30', '표 알파 10', '표 베타 20', '삭제 자료 40'],
            'after': ['베타 기준 21', '알파 기준 11', '감마 기준 31', '표 베타 21', '표 알파 11', '추가 자료 41'],
        }
        colors = {'before': ['red', 'orange', 'yellow', 'purple', 'pink', 'brown'],
                  'after': ['blue', 'green', 'cyan', 'navy', 'lime', 'white']}
        for side in ('before', 'after'):
            doc = Document()
            for index, value in enumerate(texts[side]):
                png = io.BytesIO()
                Image.new('RGB', (12, 8), colors[side][index]).save(png, format='PNG')
                raw = png.getvalue()
                self.images[f'images/{side}_{index}.png'] = raw
                if index in (3, 4):
                    table = doc.add_table(rows=2, cols=2)
                    table.cell(0, 0).merge(table.cell(0, 1)).text = '표 양식'
                    table.cell(1, 0).text = value
                    table.cell(1, 0).paragraphs[0].add_run().add_picture(io.BytesIO(raw))
                    table.cell(1, 1).text = '공통'
                    nested = table.cell(1, 1).add_table(rows=1, cols=1)
                    nested.cell(0, 0).text = '중첩표'
                else:
                    doc.add_paragraph(value)
            stream = io.BytesIO()
            doc.save(stream)
            self.docs[side] = stream.getvalue()
        self.ids = [f'change-{n:06d}' for n in range(1, 8)]
        self.items = []
        for index, (before_index, after_index) in enumerate([(0, 0), (1, 1), (2, 2), (3, 3), (4, 4), (5, None), (None, 5)], 1):
            item = dict(index=index, change_id=self.ids[index - 1], section=f'항목 {index}', page=f'전 {index} / 후 {index + 1}',
                        before=texts['before'][before_index] if before_index is not None else '',
                        after=texts['after'][after_index] if after_index is not None else '',
                        before_block=before_index, after_block=after_index,
                        before_page=index if before_index is not None else None,
                        after_page=index + 1 if after_index is not None else None,
                        context_type='table' if index in (4, 5) else 'paragraph', needs_review=False)
            for side, block in (('before', before_index), ('after', after_index)):
                if block is not None:
                    item[side + '_image'] = f'images/{side}_{block}.png'
            self.items.append(item)
        self.manifest = dict(format=g.FORMAT, schema_version=3, version='변분기 Alpha 1.1',
                             producer_version='변분기 Alpha 1.1', scope_policy=g.SCOPE_POLICY,
                             document_name='매칭 시험', item_count=len(self.items), items=self.items, warnings=[],
                             sources={side: {'sha256': hashlib.sha256(raw).hexdigest()} for side, raw in self.docs.items()})
        self.package = self.root / 'source.bdsg'
        self.write_package(self.package)
        self.analysis = g.load_package(self.package)

    def write_package(self, path):
        with zipfile.ZipFile(path, 'w') as archive:
            archive.writestr('analysis.json', json.dumps(self.manifest, ensure_ascii=False))
            for side, filename in (('before', '변경전_전체.docx'), ('after', '변경후_전체.docx')):
                archive.writestr(filename, self.docs[side])
            for name, raw in self.images.items():
                archive.writestr(name, raw)

    def rows(self):
        return g.initial_rows(self.analysis)

    def generate_rows(self, rows, name='output.docx', mode='editable'):
        output = self.root / name
        g.generate(self.analysis, rows, output, mode=mode)
        return Document(output)

    @staticmethod
    def visible_text(cell):
        return ''.join(node.text or '' for node in cell._tc.iter(qn('w:t')))

    @staticmethod
    def changed_text(cell, property_name):
        values = []
        for run in cell._tc.iter(qn('w:r')):
            props = run.find(qn('w:rPr'))
            prop = props.find(qn('w:' + property_name)) if props is not None else None
            if prop is not None and prop.get(qn('w:val'), 'true').lower() not in ('0', 'false', 'off', 'none'):
                values.extend(node.text or '' for node in run.iter(qn('w:t')))
        return ''.join(values)

    @staticmethod
    def picture_blobs(doc, cell):
        return [doc.part.rels[node.get(qn('r:embed'))].target_part.blob
                for node in cell._tc.iter(qn('a:blip')) if node.get(qn('r:embed'))]

    def test_initial_rows_have_ordered_unique_union_and_real_side_members(self):
        rows = self.rows()
        for index in range(5):
            self.assertEqual(rows[index]['ids'], [self.ids[index]])
            self.assertEqual(rows[index]['before_ids'], [self.ids[index]])
            self.assertEqual(rows[index]['after_ids'], [self.ids[index]])
        self.assertEqual(rows[5]['before_ids'], [self.ids[5]])
        self.assertEqual(rows[5]['after_ids'], [])
        self.assertEqual(rows[6]['before_ids'], [])
        self.assertEqual(rows[6]['after_ids'], [self.ids[6]])
        self.assertTrue(all(row['included'] is True for row in rows))

    def test_legacy_review_normalization_is_canonical_and_does_not_mutate_input(self):
        legacy = self.rows()
        for row in legacy:
            row.pop('before_ids'); row.pop('after_ids')
        original = copy.deepcopy(legacy)
        canonical = g.normalize_review(self.analysis, legacy)
        self.assertEqual(canonical, self.rows())
        canonical[0]['before_ids'].clear()
        self.assertEqual(legacy, original)
        self.assertEqual(self.analysis.manifest['items'], self.items)

    def test_make_review_row_derives_union_metadata_and_rejects_wrong_source_side(self):
        row = g.make_review_row(self.analysis, [self.ids[1], self.ids[0]], [self.ids[0]])
        self.assertEqual(row['before_ids'], [self.ids[1], self.ids[0]])
        self.assertEqual(row['after_ids'], [self.ids[0]])
        self.assertEqual(row['ids'], [self.ids[1], self.ids[0]])
        self.assertIn('항목 1', row['section']); self.assertIn('항목 2', row['section'])
        self.assertIn('1', row['page']); self.assertIn('2', row['page'])
        for before_ids, after_ids in [([self.ids[6]], []), ([], [self.ids[5]]), ([], []), (['unknown'], [])]:
            with self.subTest(before_ids=before_ids, after_ids=after_ids):
                with self.assertRaises(ValueError):
                    g.make_review_row(self.analysis, before_ids, after_ids)

    def test_swapped_paragraph_matches_drive_saved_review_and_character_diff(self):
        rows = self.rows()
        g.set_matching(self.analysis, rows, [0, 1], [([self.ids[0]], [self.ids[1]]), ([self.ids[1]], [self.ids[0]])])
        self.assertEqual(rows[0]['ids'], [self.ids[0], self.ids[1]])
        self.assertEqual(rows[1]['ids'], [self.ids[1], self.ids[0]])
        review = self.root / 'review.json'
        original_package = self.package.read_bytes()
        original_manifest = copy.deepcopy(self.analysis.manifest)
        g.save_review(self.analysis, rows, review)
        saved = json.loads(review.read_text(encoding='utf-8'))
        self.assertEqual(saved['format'], 'bdsg-review-v2')
        self.assertEqual(saved['package_sha256'], self.analysis.sha256)
        reloaded = g.load_review(self.analysis, review)
        self.assertEqual(reloaded, rows)
        doc = self.generate_rows(reloaded)
        table = doc.tables[0]
        self.assertIn('알파 기준 10', table.cell(1, 2).text)
        self.assertIn('알파 기준 11', table.cell(1, 3).text)
        self.assertIn('베타 기준 20', table.cell(2, 2).text)
        self.assertIn('베타 기준 21', table.cell(2, 3).text)
        self.assertEqual(self.changed_text(table.cell(1, 2), 'strike'), '0')
        self.assertEqual(self.changed_text(table.cell(1, 3), 'u'), '1')
        self.assertEqual(self.package.read_bytes(), original_package)
        self.assertEqual(self.analysis.manifest, original_manifest)

    def test_many_to_one_and_one_to_many_keep_each_source_side_once(self):
        rows = self.rows()
        g.set_matching(self.analysis, rows, [0, 1, 2],
                       [([self.ids[0], self.ids[1]], [self.ids[1]]), ([self.ids[2]], [self.ids[0], self.ids[2]])])
        self.assertEqual(len(rows), 6)
        doc = self.generate_rows(rows)
        table = doc.tables[0]
        self.assertIn('알파 기준 10', table.cell(1, 2).text)
        self.assertIn('베타 기준 20', table.cell(1, 2).text)
        self.assertIn('알파 기준 11', table.cell(1, 3).text)
        self.assertIn('감마 기준 30', table.cell(2, 2).text)
        self.assertIn('베타 기준 21', table.cell(2, 3).text)
        self.assertIn('감마 기준 31', table.cell(2, 3).text)
        full = ''.join(self.visible_text(row.cells[2]) for row in table.rows[1:])
        self.assertEqual(full.count('알파 기준 10'), 1)
        self.assertEqual(full.count('베타 기준 20'), 1)
        self.assertEqual(full.count('감마 기준 30'), 1)

    def test_separate_before_after_rows_and_reconnect_are_lossless(self):
        rows = self.rows(); original = copy.deepcopy(rows)
        g.set_matching(self.analysis, rows, [0], [([self.ids[0]], []), ([], [self.ids[0]])])
        doc = self.generate_rows(rows, 'separated.docx')
        table = doc.tables[0]
        self.assertIn('알파 기준 10', table.cell(1, 2).text)
        self.assertEqual(table.cell(1, 3).text, '(해당 없음)')
        self.assertEqual(table.cell(2, 2).text, '(해당 없음)')
        self.assertIn('베타 기준 21', table.cell(2, 3).text)
        g.set_matching(self.analysis, rows, [0, 1], [([self.ids[0]], [self.ids[0]])])
        self.assertEqual(rows, original)

    def test_merge_rows_keeps_manually_matched_side_order_and_exclusions(self):
        rows = self.rows()
        g.set_matching(self.analysis, rows, [0, 1], [([self.ids[0]], [self.ids[1]]), ([self.ids[1]], [self.ids[0]])])
        g.merge_rows(rows, [0, 1])
        self.assertEqual(rows[0]['before_ids'], [self.ids[0], self.ids[1]])
        self.assertEqual(rows[0]['after_ids'], [self.ids[1], self.ids[0]])
        self.assertEqual(rows[0]['ids'], [self.ids[0], self.ids[1]])
        rows[0]['included'] = False
        doc = self.generate_rows(rows)
        full = ''.join(self.visible_text(row.cells[2]) for row in doc.tables[0].rows[1:])
        self.assertNotIn('알파 기준 10', full)
        self.assertNotIn('베타 기준 20', full)
        self.assertIn('감마 기준 30', full)

    def test_matching_rejects_invalid_pool_or_selection_atomically(self):
        invalid = [
            ([0, 2], [([self.ids[0], self.ids[2]], [self.ids[0], self.ids[2]])]),
            ([0, 1], [([self.ids[0]], [self.ids[0]])]),
            ([0, 1], [([self.ids[0], self.ids[0], self.ids[1]], [self.ids[0], self.ids[1]])]),
            ([0], [([self.ids[0]], [self.ids[1]])]),
            ([0], [(['unknown'], [self.ids[0]])]),
            ([0], [([], [])]),
        ]
        for indices, pairs in invalid:
            with self.subTest(indices=indices, pairs=pairs):
                rows = self.rows(); snapshot = copy.deepcopy(rows)
                with self.assertRaises(ValueError): g.set_matching(self.analysis, rows, indices, pairs)
                self.assertEqual(rows, snapshot)
        rows = self.rows(); rows[0]['included'] = False; snapshot = copy.deepcopy(rows)
        with self.assertRaises(ValueError):
            g.set_matching(self.analysis, rows, [0], [([self.ids[0]], [self.ids[0]])])
        self.assertEqual(rows, snapshot)

    def test_normalization_rejects_unknown_duplicate_missing_wrong_side_empty_and_bad_union(self):
        cases = []
        rows = self.rows(); rows[0]['before_ids'] = ['unknown']; rows[0]['ids'] = ['unknown', self.ids[0]]; cases.append(rows)
        rows = self.rows(); rows[1]['before_ids'] = [self.ids[0]]; rows[1]['ids'] = [self.ids[0], self.ids[1]]; cases.append(rows)
        rows = self.rows(); rows[0]['before_ids'] = []; cases.append(rows)
        rows = self.rows(); rows[6]['before_ids'] = [self.ids[6]]; cases.append(rows)
        rows = self.rows(); rows[5]['after_ids'] = [self.ids[5]]; cases.append(rows)
        rows = self.rows(); rows[0].update(before_ids=[], after_ids=[], ids=[]); cases.append(rows)
        rows = self.rows(); rows[0]['ids'] = [self.ids[1]]; cases.append(rows)
        rows = self.rows(); rows[0]['before_ids'] = self.ids[0]; cases.append(rows)
        output = self.root / 'existing.docx'
        for index, rows in enumerate(cases):
            with self.subTest(case=index):
                original = copy.deepcopy(rows); output.write_bytes(b'previous output')
                with self.assertRaises(ValueError): g.normalize_review(self.analysis, rows)
                with self.assertRaises(ValueError): g.generate(self.analysis, rows, output)
                self.assertEqual(rows, original)
                self.assertEqual(output.read_bytes(), b'previous output')

    def test_review_v1_migrates_to_v2_and_both_remain_bound_to_package_hash(self):
        legacy = self.rows()
        for row in legacy: row.pop('before_ids'); row.pop('after_ids')
        review = self.root / 'legacy.json'
        review.write_text(json.dumps({'format': 'bdsg-review-v1', 'package_sha256': self.analysis.sha256, 'rows': legacy}), encoding='utf-8')
        loaded = g.load_review(self.analysis, review)
        self.assertEqual(loaded, self.rows())
        g.save_review(self.analysis, loaded, review)
        self.assertEqual(json.loads(review.read_text(encoding='utf-8'))['format'], 'bdsg-review-v2')
        self.manifest['document_name'] = '다른 패키지'
        other = self.root / 'other.bdsg'; self.write_package(other)
        with self.assertRaises(ValueError): g.load_review(g.load_package(other), review)

    def test_failed_review_save_preserves_existing_review_and_original_package(self):
        review = self.root / 'review.json'; review.write_bytes(b'previous review')
        original = self.package.read_bytes()
        rows = self.rows(); rows[0]['before_ids'] = []
        with self.assertRaises(ValueError): g.save_review(self.analysis, rows, review)
        self.assertEqual(review.read_bytes(), b'previous review')
        self.assertEqual(self.package.read_bytes(), original)

    def test_preview_uses_corrected_side_match_without_launching_gui(self):
        from app import App

        class Detail:
            value = ''
            state = None
            def configure(self, **kwargs): self.state = kwargs.get('state')
            def delete(self, *args): self.value = ''
            def insert(self, location, value): self.value += value

        rows = self.rows()
        g.set_matching(self.analysis, rows, [0, 1], [([self.ids[0]], [self.ids[1]]), ([self.ids[1]], [self.ids[0]])])
        detail = Detail()
        fake = SimpleNamespace(analysis=self.analysis, rows=rows, indices=lambda: [0], detail=detail)
        App.preview(fake)
        before, after = detail.value.split('[변경 후]', 1)
        self.assertIn('[변경 전]', before)
        self.assertIn('알파 기준 10', before)
        self.assertIn('알파 기준 11', after)
        self.assertNotIn('베타 기준 21', detail.value)
        self.assertEqual(detail.state, 'disabled')

    def test_reconnected_native_tables_keep_merges_nested_content_images_and_new_counterpart_diff(self):
        rows = self.rows()
        g.set_matching(self.analysis, rows, [3, 4], [([self.ids[3]], [self.ids[4]]), ([self.ids[4]], [self.ids[3]])])
        for row in rows:
            if row['before_ids'] not in ([self.ids[3]], [self.ids[4]]): row['included'] = False
        doc = self.generate_rows(rows, 'native.docx')
        table = doc.tables[0]
        old, new = table.cell(1, 2), table.cell(1, 3)
        self.assertEqual(len(list(old._tc.iter(qn('w:tbl')))), 2)
        self.assertTrue(list(old._tc.iter(qn('w:gridSpan'))))
        self.assertIn('표 알파 10', self.visible_text(old))
        self.assertIn('표 알파 11', self.visible_text(new))
        self.assertIn('중첩표', self.visible_text(old))
        self.assertEqual(self.changed_text(old, 'strike'), '0')
        self.assertEqual(self.changed_text(new, 'u'), '1')
        self.assertEqual(self.picture_blobs(doc, old), [self.images['images/before_3.png']])
        self.assertEqual(self.picture_blobs(doc, new), [self.images['images/after_4.png']])

    def test_image_mode_uses_each_selected_original_side_image(self):
        rows = self.rows()
        g.set_matching(self.analysis, rows, [0, 1], [([self.ids[0]], [self.ids[1]]), ([self.ids[1]], [self.ids[0]])])
        for row in rows[2:]: row['included'] = False
        doc = self.generate_rows(rows, 'images.docx', mode='images')
        table = doc.tables[0]
        self.assertEqual(self.picture_blobs(doc, table.cell(1, 2)), [self.images['images/before_0.png']])
        self.assertEqual(self.picture_blobs(doc, table.cell(1, 3)), [self.images['images/after_1.png']])
        self.assertEqual(self.picture_blobs(doc, table.cell(2, 2)), [self.images['images/before_1.png']])
        self.assertEqual(self.picture_blobs(doc, table.cell(2, 3)), [self.images['images/after_0.png']])

    def test_cancel_with_manual_matching_preserves_existing_output_and_review_rows(self):
        rows = self.rows()
        g.set_matching(self.analysis, rows, [0, 1], [([self.ids[0]], [self.ids[1]]), ([self.ids[1]], [self.ids[0]])])
        snapshot = copy.deepcopy(rows)
        output = self.root / 'previous.docx'; output.write_bytes(b'previous output')
        cancelled = threading.Event()
        with self.assertRaises(g.Cancelled):
            g.generate(self.analysis, rows, output, cancel=cancelled, progress=lambda n, total: cancelled.set())
        self.assertEqual(output.read_bytes(), b'previous output')
        self.assertEqual(rows, snapshot)


if __name__ == '__main__':
    unittest.main()
