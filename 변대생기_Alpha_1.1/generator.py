"""Stage two: trusted local analysis package -> reviewable Word comparison."""
from __future__ import annotations

import copy
import difflib
import hashlib
import io
import json
import os
import re
import tempfile
import threading
import zipfile
from dataclasses import dataclass
from itertools import zip_longest
from pathlib import Path, PurePosixPath

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

VERSION = "변대생기 Alpha 1.1"
SCOPE_POLICY = 'front_matter_and_body_excluding_toc_and_appendices_v1'
FORMAT = "byeondaesaenggi-analysis-v1"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MAX_PACKAGE = 512 * 1024 * 1024


class Cancelled(Exception):
    pass


def check_cancel(cancel):
    if cancel is not None and cancel.is_set():
        raise Cancelled("작업을 취소했습니다.")


def clean_text(value):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", value)


@dataclass
class Analysis:
    path: Path
    sha256: str
    manifest: dict
    before_doc: object
    after_doc: object
    images: dict


def load_package(path, cancel=None):
    path = Path(path).resolve()
    if path.suffix.lower() != ".bdsg":
        raise ValueError("변분기 Alpha 1.1, 1.0, 0.5 또는 0.4에서 생성한 .bdsg 파일을 선택하십시오.")
    if path.stat().st_size > MAX_PACKAGE:
        raise ValueError("분석 패키지가 512MB 제한을 초과합니다.")
    raw = path.read_bytes()
    check_cancel(cancel)
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise ValueError("분석 패키지에 중복 파일명이 있습니다.")
        if sum(i.file_size for i in archive.infolist()) > MAX_PACKAGE:
            raise ValueError("분석 패키지 압축 해제 크기가 제한을 초과합니다.")
        for name in names:
            p = PurePosixPath(name)
            if p.is_absolute() or ".." in p.parts or "\\" in name or ":" in name:
                raise ValueError("분석 패키지 내부 경로가 올바르지 않습니다.")
        m = json.loads(archive.read("analysis.json"))
        supported = (m.get('version'), m.get('schema_version')) in {('변분기 Alpha 0.4', 2), ('변분기 Alpha 0.5', 3), ('변분기 Alpha 1.0', 3), ('변분기 Alpha 1.1', 3)}
        if m.get("format") != FORMAT or not supported:
            raise ValueError("변분기 Alpha 1.1·1.0·0.5(schema_version 3) 또는 0.4(schema_version 2) 패키지만 지원합니다.")
        if m.get('schema_version') == 3 and m.get('scope_policy') != SCOPE_POLICY:
            raise ValueError('Alpha 1.1·1.0·0.5 패키지의 비교 범위 정책이 올바르지 않습니다.')
        items = m.get("items")
        if not isinstance(items, list) or m.get("item_count") != len(items):
            raise ValueError("변경 항목 수가 올바르지 않습니다.")
        if not isinstance(m.get("document_name"), str) or not isinstance(m.get("warnings", []), list):
            raise ValueError("문서명 또는 경고 정보가 올바르지 않습니다.")
        if any(not isinstance(x, str) for x in m.get("warnings", [])):
            raise ValueError("경고 정보가 올바르지 않습니다.")
        docs = {}
        for side, name in (("before", "변경전_전체.docx"), ("after", "변경후_전체.docx")):
            check_cancel(cancel)
            content = archive.read(name)
            expected = m.get("sources", {}).get(side, {}).get("sha256")
            if not expected or hashlib.sha256(content).hexdigest() != expected:
                raise ValueError(f"{side} 원문 해시가 분석 결과와 다릅니다.")
            docs[side] = Document(io.BytesIO(content))
        ids = set()
        images = {}
        for item in items:
            check_cancel(cancel)
            if not isinstance(item, dict) or not isinstance(item.get("change_id"), str) or item["change_id"] in ids:
                raise ValueError("변경 ID가 누락되거나 중복되었습니다.")
            ids.add(item["change_id"])
            for key in ("section", "before", "after", "page"):
                if not isinstance(item.get(key), str):
                    raise ValueError(f"변경 항목의 {key} 값이 올바르지 않습니다.")
            for side in ("before", "after"):
                index = item.get(side + "_block")
                if index is not None and (type(index) is not int or index < 0 or index >= len(docs[side]._element.body) or docs[side]._element.body[index].tag == qn("w:sectPr")):
                    raise ValueError(f"{item['change_id']}: {side} 본문 위치가 올바르지 않습니다.")
                if index is None and item[side]:
                    raise ValueError("본문 위치가 없는 항목에 원문이 있습니다.")
                image_name = item.get(side + "_image")
                if image_name:
                    if not isinstance(image_name, str) or not image_name.startswith("images/"):
                        raise ValueError("이미지 경로가 올바르지 않습니다.")
                    if image_name in names:
                        images[image_name] = archive.read(image_name)
                    else:
                        item["needs_review"] = True
                        item["image_error"] = str(item.get("image_error", "")) + f"\n이미지 누락: {image_name}"
        return Analysis(path, hashlib.sha256(raw).hexdigest(), m, docs["before"], docs["after"], images)


def initial_rows(analysis):
    return [{"ids": [item["change_id"]],
             "before_ids": [item["change_id"]] if item.get('before_block') is not None else [],
             "after_ids": [item["change_id"]] if item.get('after_block') is not None else [],
             "section": item["section"], "page": item["page"], "included": True}
            for item in analysis.manifest["items"]]


def reanalysis_warnings(manifest):
    if manifest.get('schema_version') == 2:
        return ['기존 Alpha 0.4 분석 패키지입니다. 문서 앞부분 변경을 포함하려면 변분기 Alpha 1.1에서 원문을 다시 분석하십시오.']
    affected = []
    for side, label in (('before', '변경 전'), ('after', '변경 후')):
        scope = manifest.get('scope_' + side)
        if not isinstance(scope, dict):
            continue
        end, front_end = scope.get('paragraph_end'), scope.get('front_matter_end')
        if type(end) is int and type(front_end) is int and end < front_end:
            affected.append(label)
    if affected:
        return [f"{'·'.join(affected)} 분석 범위가 목차 종료 전에서 끝났습니다. 본문 변경이 누락되었을 수 있으므로 변분기 Alpha 1.1에서 원문을 다시 분석하십시오. 변대생기는 기존 패키지에 없는 변경을 복원할 수 없습니다."]
    return []


def merge_rows(rows, indices):
    selected = _selected_indices(rows, indices)
    if len(selected) < 2 or selected != list(range(selected[0], selected[-1] + 1)):
        raise ValueError("연속된 항목을 2개 이상 선택하십시오.")
    if any(not rows[i]["included"] for i in selected):
        raise ValueError("제외된 항목은 포함으로 복원한 뒤 병합하십시오.")
    canonical = ['before_ids' in rows[i] or 'after_ids' in rows[i] for i in selected]
    if any(canonical) and not all(canonical):
        raise ValueError("이전 검토 형식은 normalize_review로 변환한 뒤 병합하십시오.")
    merged = {"ids": list(dict.fromkeys(cid for i in selected for cid in rows[i]["ids"])),
              "section": " / ".join(dict.fromkeys(rows[i]["section"] for i in selected)),
              "page": "\n".join(dict.fromkeys(rows[i]["page"] for i in selected)), "included": True}
    if all(canonical):
        for side in ('before', 'after'):
            if any(not isinstance(rows[i].get(side + '_ids'), list) for i in selected):
                raise ValueError("전후 매칭 ID 목록이 올바르지 않습니다.")
            merged[side + '_ids'] = [cid for i in selected for cid in rows[i][side + '_ids']]
        merged['ids'] = list(dict.fromkeys(merged['before_ids'] + merged['after_ids']))
    rows[selected[0]:selected[-1] + 1] = [merged]


def save_review(analysis, rows, path):
    path = Path(path).resolve()
    if path.suffix.lower() != ".json" or path == analysis.path:
        raise ValueError("검토 결과는 별도의 .json 파일에 저장하십시오.")
    normalized = normalize_review(analysis, rows)
    data = {"format": "bdsg-review-v2", "package_sha256": analysis.sha256, "rows": normalized}
    atomic_write(path, lambda tmp: tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"))


def validate_rows(analysis, rows):
    normalize_review(analysis, rows)


def normalize_review(analysis, rows):
    """Validate complete side coverage, including excluded rows, and copy rows."""
    by_id = {i['change_id']: i for i in analysis.manifest['items']}
    known = set(by_id)
    if not isinstance(rows, list):
        raise ValueError("검토 항목 형식이 올바르지 않습니다.")
    legacy_seen = set()
    side_seen = {'before': set(), 'after': set()}
    modes = set()
    normalized = []
    for row in rows:
        if not isinstance(row, dict) or not isinstance(row.get("ids"), list) or not row["ids"] or any(not isinstance(cid, str) for cid in row["ids"]):
            raise ValueError("검토 변경 ID가 올바르지 않습니다.")
        if any(cid not in known for cid in row['ids']) or len(row['ids']) != len(set(row['ids'])):
            raise ValueError("검토 항목에 알 수 없거나 중복된 ID가 있습니다.")
        if type(row.get("included")) is not bool or any(not isinstance(row.get(k), str) for k in ("section", "page")):
            raise ValueError("검토 정보가 올바르지 않습니다.")
        has_before, has_after = 'before_ids' in row, 'after_ids' in row
        if has_before != has_after:
            raise ValueError("전후 매칭 ID 목록이 모두 필요합니다.")
        modes.add(has_before)
        if len(modes) > 1:
            raise ValueError("이전 검토와 새 매칭 형식을 한 목록에 섞을 수 없습니다.")
        if has_before:
            before_ids, after_ids = row['before_ids'], row['after_ids']
        else:
            if any(cid in legacy_seen for cid in row['ids']):
                raise ValueError("검토 항목에 중복된 ID가 있습니다.")
            legacy_seen.update(row['ids'])
            before_ids = [cid for cid in row['ids'] if by_id[cid].get('before_block') is not None]
            after_ids = [cid for cid in row['ids'] if by_id[cid].get('after_block') is not None]
        for side, ids in (('before', before_ids), ('after', after_ids)):
            _validate_side_ids(by_id, side, ids)
            if any(cid in side_seen[side] for cid in ids):
                raise ValueError(f"{side} 검토 항목에 중복된 ID가 있습니다.")
            side_seen[side].update(ids)
        union = list(dict.fromkeys(before_ids + after_ids))
        if not union:
            raise ValueError("변경 전·후가 모두 없는 검토 행은 사용할 수 없습니다.")
        if has_before and row['ids'] != union:
            raise ValueError("검토 ids가 전후 매칭 ID 합집합과 다릅니다.")
        normalized.append({'ids': union, 'before_ids': list(before_ids), 'after_ids': list(after_ids),
                           'section': row['section'], 'page': row['page'], 'included': row['included']})
    if modes == {False} and legacy_seen != known:
        raise ValueError("검토 파일에 누락된 변경 ID가 있습니다.")
    for side in ('before', 'after'):
        expected = {cid for cid, item in by_id.items() if item.get(side + '_block') is not None}
        if side_seen[side] != expected:
            raise ValueError(f"검토 파일에 누락된 {side} 변경 ID가 있습니다.")
    return normalized


def _validate_side_ids(by_id, side, ids):
    if not isinstance(ids, list) or any(not isinstance(cid, str) for cid in ids):
        raise ValueError(f"{side} 매칭 ID 목록이 올바르지 않습니다.")
    if len(ids) != len(set(ids)):
        raise ValueError(f"{side} 매칭 ID가 중복되었습니다.")
    if any(cid not in by_id or by_id[cid].get(side + '_block') is None for cid in ids):
        raise ValueError(f"{side}에 존재하지 않는 원문 ID입니다.")


def make_review_row(analysis, before_ids, after_ids, included=True):
    by_id = {i['change_id']: i for i in analysis.manifest['items']}
    if type(included) is not bool:
        raise ValueError("검토 포함 여부가 올바르지 않습니다.")
    _validate_side_ids(by_id, 'before', before_ids)
    _validate_side_ids(by_id, 'after', after_ids)
    union = list(dict.fromkeys(before_ids + after_ids))
    if not union:
        raise ValueError("변경 전·후가 모두 없는 검토 행은 사용할 수 없습니다.")
    page_parts = []
    for side, ids, label in (('before', before_ids, '전'), ('after', after_ids, '후')):
        values = []
        for cid in ids:
            item = by_id[cid]
            value = item.get(side + '_page')
            values.append(str(value) if value is not None and str(value) else '원분석 page: ' + item['page'])
        page_parts.append(label + ' ' + (', '.join(dict.fromkeys(values)) if ids else '없음'))
    return {'ids': union, 'before_ids': list(before_ids), 'after_ids': list(after_ids),
            'section': ' / '.join(dict.fromkeys(by_id[cid]['section'] for cid in union)),
            'page': ' / '.join(page_parts), 'included': included}


def _selected_indices(rows, indices):
    if not isinstance(indices, (list, tuple)) or any(type(i) is not int or i < 0 or i >= len(rows) for i in indices):
        raise ValueError("선택한 검토 행이 올바르지 않습니다.")
    selected = sorted(set(indices))
    if not selected or selected != list(range(selected[0], selected[-1] + 1)):
        raise ValueError("연속된 검토 행을 선택하십시오.")
    return selected


def set_matching(analysis, rows, indices, pairs):
    """Replace a selected side pool only after complete review validation."""
    canonical = normalize_review(analysis, rows)
    selected = _selected_indices(canonical, indices)
    if any(not canonical[i]['included'] for i in selected):
        raise ValueError("제외된 항목은 포함으로 복원한 뒤 매칭을 수정하십시오.")
    if not isinstance(pairs, (list, tuple)) or not pairs:
        raise ValueError("새 전후 매칭 행을 하나 이상 지정하십시오.")
    replacements = []
    for pair in pairs:
        if not isinstance(pair, (list, tuple)) or len(pair) != 2:
            raise ValueError("새 매칭은 변경 전·후 ID 목록 한 쌍이어야 합니다.")
        replacements.append(make_review_row(analysis, pair[0], pair[1]))
    for side in ('before', 'after'):
        pool = [cid for i in selected for cid in canonical[i][side + '_ids']]
        new_pool = [cid for row in replacements for cid in row[side + '_ids']]
        if len(new_pool) != len(set(new_pool)) or set(new_pool) != set(pool):
            raise ValueError(f"선택한 {side} 원문 ID를 빠짐·중복 없이 새 매칭에 배정하십시오.")
    candidate = canonical[:selected[0]] + replacements + canonical[selected[-1] + 1:]
    candidate = normalize_review(analysis, candidate)
    rows[:] = candidate
    return rows


def load_review(analysis, path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(data, dict) or data.get("format") not in ("bdsg-review-v1", "bdsg-review-v2") or data.get("package_sha256") != analysis.sha256:
        raise ValueError("이 분석 패키지에서 저장한 검토 파일이 아닙니다.")
    rows = data.get('rows')
    if not isinstance(rows, list):
        raise ValueError("검토 항목 형식이 올바르지 않습니다.")
    canonical = data['format'] == 'bdsg-review-v2'
    if any(not isinstance(row, dict) or ('before_ids' in row) != canonical or ('after_ids' in row) != canonical for row in rows):
        raise ValueError("검토 파일 버전과 전후 매칭 정보가 다릅니다.")
    return normalize_review(analysis, rows)


def atomic_write(path, writer):
    path = Path(path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(dir=path.parent, suffix=path.suffix)
    os.close(fd)
    tmp = Path(name)
    try:
        writer(tmp)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


def set_font(run, size=8, changed=None):
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string("000000")
    run.font.strike = False
    run.font.underline = False
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), "맑은 고딕")
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attr), None)
    complex_size = rpr.find(qn("w:szCs"))
    if complex_size is None:
        complex_size = OxmlElement("w:szCs")
        rpr.append(complex_size)
    complex_size.set(qn("w:val"), str(size * 2))
    if changed:
        run.font.color.rgb = RGBColor.from_string("C00000" if changed == "before" else "0000FF")
        run.font.strike = changed == "before"
        run.font.underline = changed == "after"


def add_diff(cell, old, new, side):
    text = old if side == "before" else new
    p = cell.add_paragraph() if cell.paragraphs[0].text or len(cell._tc) > 2 else cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    for tag, a, b, c, d in difflib.SequenceMatcher(None, old, new, autojunk=False).get_opcodes():
        value = text[a:b] if side == "before" else text[c:d]
        if value:
            run = p.add_run(clean_text(value))
            set_font(run, changed=side if tag != "equal" else None)
    if not text:
        set_font(p.add_run("(해당 없음)"))


def copy_table(cell, source_doc, index, other_doc, other_index, output_doc, side):
    """Retain merged/nested native tables; import images and external links only."""
    node = copy.deepcopy(source_doc._element.body[index])
    if node.tag != qn("w:tbl"):
        raise ValueError("대응 블록이 일반 표가 아닙니다")
    unsupported = ("object", "altChunk", "sdt", "fldChar", "instrText", "fldSimple")
    if any(list(node.iter(qn("w:" + tag))) for tag in unsupported):
        raise ValueError("특수한 표 요소가 있어 원문 텍스트를 사용합니다")
    for element in node.iter():
        for attr, value in list(element.attrib.items()):
            if attr.startswith("{" + R + "}"):
                rel = source_doc.part.rels[value]
                if rel.is_external:
                    element.set(attr, output_doc.part.relate_to(rel.target_ref, rel.reltype, is_external=True))
                elif rel.reltype.endswith("/image"):
                    rid, _ = output_doc.part.get_or_add_image(io.BytesIO(rel.target_part.blob))
                    element.set(attr, rid)
                else:
                    raise ValueError("이미지 이외의 포함 개체는 텍스트로 대체합니다")
    # Compare the complete visible table text. Run boundaries do not affect diff.
    own = "".join(t.text or "" for t in node.iter(qn("w:t")))
    other = ""
    if other_index is not None and other_doc._element.body[other_index].tag == qn('w:tbl'):
        other = "".join(t.text or "" for t in other_doc._element.body[other_index].iter(qn("w:t")))
    spans = [(a, b) for tag, a, b, _, _ in difflib.SequenceMatcher(None, own, other, autojunk=False).get_opcodes() if tag != "equal"]
    offset = 0
    from docx.text.run import Run
    for text_node in list(node.iter(qn("w:t"))):
        text = text_node.text or ""
        parent = text_node.getparent()
        if parent.tag != qn("w:r"):
            raise ValueError("복잡한 표 텍스트는 원문으로 대체합니다")
        boundaries = {0, len(text)}
        for a, b in spans:
            boundaries.update(max(0, min(len(text), x - offset)) for x in (a, b))
        positions = sorted(boundaries)
        for start, end in zip(positions, positions[1:]):
            if start == end:
                continue
            run_node = OxmlElement("w:r")
            props = parent.find(qn("w:rPr"))
            if props is not None:
                run_node.append(copy.deepcopy(props))
            t = OxmlElement("w:t"); t.text = text[start:end]; t.set(qn("xml:space"), "preserve")
            run_node.append(t)
            parent.addprevious(run_node)
            changed = any(a <= offset + start < b for a, b in spans)
            set_font(Run(run_node, cell.paragraphs[0]), size=7, changed=side if changed else None)
        parent.remove(text_node)
        offset += len(text)
    # Avoid source IDs, styles, automatic numbering and page breaks leaking into output.
    for tag in ("bookmarkStart", "bookmarkEnd", "commentRangeStart", "commentRangeEnd", "commentReference", "pStyle", "rStyle", "numPr", "sectPr", "pageBreakBefore", "keepNext", "cantSplit", "trHeight", "noWrap", "tblInd"):
        for element in list(node.iter(qn("w:" + tag))):
            element.getparent().remove(element)
    notes = []
    from docx.oxml import parse_xml
    for tag, part_type, label in (("footnoteReference", "/footnotes", "각주"), ("endnoteReference", "/endnotes", "미주")):
        refs = list(node.iter(qn("w:" + tag)))
        if not refs:
            continue
        rel = next((r for r in source_doc.part.rels.values() if r.reltype.endswith(part_type)), None)
        if rel is None:
            raise ValueError(f"{label} 원문을 찾을 수 없습니다")
        note_root = parse_xml(rel.target_part.blob)
        values = {n.get(qn("w:id")): "".join(t.text or "" for t in n.iter(qn("w:t"))) for n in note_root}
        for ref in refs:
            note_id = ref.get(qn("w:id"))
            if note_id not in values:
                raise ValueError(f"{label} {note_id} 원문을 찾을 수 없습니다")
            t = OxmlElement("w:t"); t.text = f"[{label} {note_id}]"
            ref.getparent().replace(ref, t)
            notes.append(f"[{label} {note_id}] {values[note_id]}")
    for table_node in node.iter(qn("w:tbl")):
        grid = table_node.find(qn("w:tblGrid"))
        cols = list(grid) if grid is not None else []
        widths = [max(1, int(col.get(qn("w:w"), "1"))) for col in cols]
        factor = min(1, 5100 / max(1, sum(widths)))
        for col, width in zip(cols, widths):
            col.set(qn("w:w"), str(max(1, round(width * factor))))
        for el in table_node.iter(qn("w:tcW")):
            if el.get(qn("w:type")) == "dxa":
                el.set(qn("w:w"), str(max(1, round(int(el.get(qn("w:w"), "1")) * factor))))
        for el in table_node.iter(qn("w:tblW")):
            el.set(qn("w:type"), "pct"); el.set(qn("w:w"), "5000")
    for p in node.iter(qn("w:p")):
        from docx.text.paragraph import Paragraph
        paragraph = Paragraph(p, cell)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
    for ext in node.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"):
        cx = int(ext.get("cx", "1")); cy = int(ext.get("cy", "1"))
        if cx > Cm(9).emu:
            ext.set("cx", str(Cm(9).emu)); ext.set("cy", str(round(cy * Cm(9).emu / cx)))
    cell._tc.insert(len(cell._tc) - 1, node)
    for note in dict.fromkeys(notes):
        p = cell.add_paragraph()
        set_font(p.add_run(clean_text(note)), size=7)
    return bool(notes)


def generate(analysis, rows, output, *, mode="editable", document_name=None, cancel=None, progress=None):
    output = Path(output).resolve()
    if output.suffix.lower() != ".docx" or output == analysis.path:
        raise ValueError("출력 경로는 별도의 .docx 파일이어야 합니다.")
    if mode not in ("editable", "images"):
        raise ValueError("출력 방식이 올바르지 않습니다.")
    canonical = normalize_review(analysis, rows)
    selected = [r for r in canonical if r["included"]]
    if not selected:
        raise ValueError("생성할 항목이 없습니다. 변경이 없거나 모든 항목이 제외되었습니다.")
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    sec.left_margin = sec.right_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name = "맑은 고딕"; normal.font.size = Pt(8)
    normal.paragraph_format.space_after = Pt(3)
    title_style = doc.styles["Title"]
    title_style.font.color.rgb = RGBColor.from_string("000000")
    for border in list(title_style._element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    title = doc.add_paragraph(style="Title")
    set_font(title.add_run(clean_text(document_name or analysis.manifest["document_name"])), size=11)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    widths = [2.8, 2.5, 9.5, 9.5, 2.4]
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    header = table.rows[0]
    repeat = OxmlElement("w:tblHeader"); header._tr.get_or_add_trPr().append(repeat)
    for cell, label, width in zip(header.cells, ["Section", "page", "변경 전", "변경 후", "변경 사유"], widths):
        cell.width = Cm(width); cell.text = label
        for run in cell.paragraphs[0].runs:
            set_font(run); run.bold = True
        shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E7E6E6"); cell._tc.get_or_add_tcPr().append(shade)
    by_id = {item["change_id"]: item for item in analysis.manifest["items"]}
    warnings = list(analysis.manifest.get("warnings", []))
    warnings.extend(reanalysis_warnings(analysis.manifest))
    for number, review in enumerate(selected, 1):
        check_cancel(cancel)
        if progress:
            progress(number, len(selected))
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = Cm(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cells[0].text = clean_text(review["section"]); cells[1].text = clean_text(review["page"])
        for cid in review["ids"]:
            source_item = by_id[cid]
            if source_item.get("needs_review") or source_item.get("image_error"):
                warnings.append(f"{cid}: 대응 관계·페이지·캡처 확인 필요. {source_item.get('image_error', '')}")
        virtual_pairs = list(zip_longest(review['before_ids'], review['after_ids']))
        changed_matching = any(
            old_id != new_id and not (
                old_id is None and by_id[new_id].get('before_block') is None or
                new_id is None and by_id[old_id].get('after_block') is None)
            for old_id, new_id in virtual_pairs)
        if changed_matching:
            warnings.append('수동 매칭 변경: 변경 전 [' + ', '.join(review['before_ids']) +
                            '] / 변경 후 [' + ', '.join(review['after_ids']) + '] 원문 ID를 확인하십시오.')
        for old_id, new_id in virtual_pairs:
            source_ids = {'before': old_id, 'after': new_id}
            item = {}
            for side, source_id in source_ids.items():
                source_item = by_id[source_id] if source_id is not None else {}
                item[side] = source_item.get(side, '')
                item[side + '_block'] = source_item.get(side + '_block')
                item[side + '_image'] = source_item.get(side + '_image')
            for side, cell in (("before", cells[2]), ("after", cells[3])):
                check_cancel(cancel)
                cid = source_ids[side]
                if cid is None:
                    continue
                idx = item.get(side + "_block")
                src = analysis.before_doc if side == "before" else analysis.after_doc
                other = analysis.after_doc if side == "before" else analysis.before_doc
                if mode == "images" and item.get(side + "_image") in analysis.images:
                    try:
                        p = cell.add_paragraph()
                        run = p.add_run()
                        shape = run.add_picture(io.BytesIO(analysis.images[item[side + "_image"]]))
                        scale = min(Cm(9).emu / shape.width, Cm(15.5).emu / shape.height, 1)
                        shape.width = round(shape.width * scale); shape.height = round(shape.height * scale)
                        continue
                    except Exception as exc:
                        warnings.append(f"{cid} {side}: 이미지 삽입 실패, 텍스트 대체 ({type(exc).__name__})")
                if mode == "editable" and idx is not None and src._element.body[idx].tag == qn("w:tbl"):
                    try:
                        has_notes = copy_table(cell, src, idx, other, item.get(("after" if side == "before" else "before") + "_block"), doc, side)
                        if has_notes:
                            warnings.append(f"{cid} {side}: 원문 표의 각주·미주를 표 아래 텍스트로 옮겼습니다. 주 번호는 원본 XML ID이며 표시 번호와 다를 수 있습니다.")
                        continue
                    except Exception as exc:
                        warnings.append(f"{cid} {side}: 원문 표를 텍스트로 대체 ({type(exc).__name__}: {exc})")
                if mode == "images" and idx is not None:
                    warnings.append(f"{cid} {side}: 이미지 없음. 원문 텍스트로 대체")
                if mode == "editable" and idx is not None:
                    node = src._element.body[idx]
                    if any(list(node.iter(qn("w:" + tag))) for tag in ("drawing", "pict", "object")) or list(node.iter("{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")):
                        warnings.append(f"{cid} {side}: 텍스트 이외의 그림·개체·수식은 이미지 방식 또는 원문 확인 필요")
                add_diff(cell, item["before"], item["after"], side)
        for side, cell in (('before', cells[2]), ('after', cells[3])):
            if not review[side + '_ids']:
                add_diff(cell, '', '', side)
        for cell in (cells[0], cells[1]):
            for p in cell.paragraphs:
                for run in p.runs:
                    set_font(run)
    # Drawing IDs must be unique even after copying multiple original tables.
    for n, el in enumerate(doc._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"), 1):
        el.set("id", str(n))
    warnings = list(dict.fromkeys(warnings))
    if warnings:
        doc.add_paragraph("확인사항")
        for warning in warnings:
            p = doc.add_paragraph()
            set_font(p.add_run(clean_text(warning)), size=8)
    check_cancel(cancel)
    def write(tmp):
        doc.save(tmp)
        check_cancel(cancel)
    atomic_write(output, write)
    return warnings
