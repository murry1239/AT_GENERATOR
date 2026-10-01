import copy
import hashlib
import io
import json
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from generator import Cancelled, FORMAT, generate, initial_rows, load_package, load_review, merge_rows, save_review


class GeneratorTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.docs = {}
        for side, text in (("before", "기존 용량 10 mg"), ("after", "변경 용량 20 mg")):
            doc = Document(); doc.add_paragraph(text)
            table = doc.add_table(rows=2, cols=2); table.cell(0, 0).merge(table.cell(0, 1)).text = "용량 정보"
            table.cell(1, 0).text = text; table.cell(1, 1).text = "공통"
            nested = table.cell(1, 1).add_table(rows=1, cols=1); nested.cell(0, 0).text = "중첩표"
            data = io.BytesIO(); doc.save(data); self.docs[side] = data.getvalue()
        self.items = [dict(index=1, change_id="change-000001", section="1. 용량", page="전 1 / 후 2", before="기존 용량 10 mg", after="변경 용량 20 mg", before_block=0, after_block=0, context_type="paragraph", needs_review=False),
                      dict(index=2, change_id="change-000002", section="2. 표", page="전 1 / 후 2", before="기존 용량 10 mg\n공통", after="변경 용량 20 mg\n공통", before_block=1, after_block=1, context_type="table", needs_review=False)]
        self.manifest = dict(format=FORMAT, schema_version=2, version="변분기 Alpha 0.4", document_name="시험 문서", item_count=2, items=self.items, warnings=[], sources={s: {"sha256": hashlib.sha256(raw).hexdigest()} for s, raw in self.docs.items()})
        self.package = self.root / "analysis.bdsg"; self.write_package()

    def write_package(self, extra=None):
        with zipfile.ZipFile(self.package, "w") as z:
            z.writestr("analysis.json", json.dumps(self.manifest, ensure_ascii=False))
            for side, name in (("before", "변경전_전체.docx"), ("after", "변경후_전체.docx")): z.writestr(name, self.docs[side])
            if extra: z.writestr(*extra)

    def test_real_docx_contains_five_columns_native_merged_and_nested_tables_and_diff(self):
        a = load_package(self.package); out = self.root / "out.docx"
        original = self.package.read_bytes()
        self.assertEqual(generate(a, initial_rows(a), out), [])
        d = Document(out); table = d.tables[0]
        self.assertEqual([c.text for c in table.rows[0].cells], ["Section", "page", "변경 전", "변경 후", "변경 사유"])
        self.assertEqual(len(table.rows), 3)
        self.assertTrue(all(not row.cells[4].text for row in table.rows[1:]))
        self.assertGreater(d.sections[0].page_width, d.sections[0].page_height)
        self.assertTrue(table.rows[0]._tr.find("w:trPr/w:tblHeader", table.rows[0]._tr.nsmap) is not None)
        self.assertEqual(len(table.cell(2, 2).tables), 1)
        self.assertTrue(list(table.cell(2, 2)._tc.iter(qn("w:gridSpan"))))
        self.assertEqual(len(list(table.cell(2, 2)._tc.iter(qn("w:tbl")))), 2)
        old = table.cell(1, 2).paragraphs[0].runs; new = table.cell(1, 3).paragraphs[0].runs
        self.assertTrue(any(r.font.strike and str(r.font.color.rgb) == "C00000" for r in old))
        self.assertTrue(any(r.font.underline and str(r.font.color.rgb) == "0000FF" for r in new))
        self.assertTrue(list(table.cell(2, 2)._tc.iter(qn("w:strike"))))
        self.assertEqual(self.package.read_bytes(), original)

    def test_merge_review_roundtrip_and_exclusion(self):
        a = load_package(self.package); rows = initial_rows(a)
        merge_rows(rows, [0, 1]); self.assertEqual(len(rows), 1); self.assertEqual(len(rows[0]["ids"]), 2)
        review = self.root / "review.json"; save_review(a, rows, review); self.assertEqual(load_review(a, review), rows)
        generate(a, rows, self.root / "merged.docx")
        self.assertEqual(len(Document(self.root / "merged.docx").tables[0].rows), 2)
        rows[0]["included"] = False
        with self.assertRaisesRegex(ValueError, "생성할 항목"): generate(a, rows, self.root / "out.docx")

    def test_missing_image_falls_back_without_dropping_text(self):
        self.items[0]["before_image"] = "images/missing.png"; self.write_package()
        a = load_package(self.package); out = self.root / "out.docx"
        warnings = generate(a, initial_rows(a), out, mode="images")
        self.assertTrue(warnings)
        self.assertIn("기존 용량 10 mg", Document(out).tables[0].cell(1, 2).text)

    def test_rejects_traversal_hash_mismatch_invalid_block_and_schema(self):
        self.write_package(("../escape", b"bad"))
        with self.assertRaises(ValueError): load_package(self.package)
        self.manifest["sources"]["before"]["sha256"] = "0" * 64; self.write_package()
        with self.assertRaisesRegex(ValueError, "해시"): load_package(self.package)
        self.manifest["sources"]["before"]["sha256"] = hashlib.sha256(self.docs["before"]).hexdigest()
        self.items[0]["before_block"] = -1; self.write_package()
        with self.assertRaisesRegex(ValueError, "본문 위치"): load_package(self.package)
        self.items[0]["before_block"] = 0; self.manifest["schema_version"] = 1; self.write_package()
        with self.assertRaisesRegex(ValueError, "schema_version"): load_package(self.package)

    def test_review_cannot_silently_drop_or_duplicate_changes(self):
        a = load_package(self.package); rows = initial_rows(a)
        with self.assertRaises(ValueError): generate(a, rows[:1], self.root / "out.docx")
        rows[1]["ids"] = rows[0]["ids"]
        with self.assertRaises(ValueError): generate(a, rows, self.root / "out.docx")

    def test_cancel_keeps_existing_output(self):
        a = load_package(self.package); out = self.root / "out.docx"; out.write_bytes(b"previous")
        cancel = threading.Event()
        def progress(n, total): cancel.set()
        with self.assertRaises(Cancelled): generate(a, initial_rows(a), out, cancel=cancel, progress=progress)
        self.assertEqual(out.read_bytes(), b"previous")

    def test_review_save_cannot_overwrite_package(self):
        a = load_package(self.package); original = self.package.read_bytes()
        with self.assertRaises(ValueError): save_review(a, initial_rows(a), self.package)
        self.assertEqual(self.package.read_bytes(), original)

    def test_insert_delete_and_unchanged_text_are_not_lost(self):
        self.items[0].update(before="", before_block=None)
        self.items[1].update(after="", after_block=None)
        self.write_package(); a = load_package(self.package); out = self.root / "out.docx"
        generate(a, initial_rows(a), out)
        d = Document(out); t = d.tables[0]
        self.assertEqual(t.cell(1, 2).text, "(해당 없음)")
        self.assertIn("변경 용량 20 mg", t.cell(1, 3).text)
        self.assertEqual(t.cell(2, 3).text, "(해당 없음)")
        self.assertIn("기존 용량 10 mg", t.cell(2, 2).tables[0].cell(1, 0).text)

    def test_native_table_image_relationships_are_imported(self):
        from PIL import Image
        for side in ("before", "after"):
            d = Document(io.BytesIO(self.docs[side]))
            png = io.BytesIO(); Image.new("RGB", (80, 30), "red" if side == "before" else "blue").save(png, format="PNG"); png.seek(0)
            d.tables[0].cell(1, 0).paragraphs[0].add_run().add_picture(png)
            data = io.BytesIO(); d.save(data); self.docs[side] = data.getvalue()
            self.manifest["sources"][side]["sha256"] = hashlib.sha256(self.docs[side]).hexdigest()
        self.write_package(); a = load_package(self.package); out = self.root / "out.docx"
        generate(a, initial_rows(a), out)
        d = Document(out)
        self.assertEqual(len(d.inline_shapes), 2)
        for shape in d.inline_shapes:
            rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            self.assertTrue(d.part.rels[rid].target_part.blob.startswith(b"\x89PNG"))

    def test_native_table_footnote_content_is_preserved(self):
        from lxml import etree as ET
        from docx.oxml import OxmlElement
        d = Document(io.BytesIO(self.docs["before"]))
        ref = OxmlElement("w:footnoteReference"); ref.set(qn("w:id"), "1")
        d.tables[0].cell(1, 0).paragraphs[0].add_run()._r.append(ref)
        data = io.BytesIO(); d.save(data)
        with zipfile.ZipFile(io.BytesIO(data.getvalue())) as z: parts = {n:z.read(n) for n in z.namelist()}
        rels = ET.fromstring(parts['word/_rels/document.xml.rels'])
        ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        ET.SubElement(rels, '{'+ns+'}Relationship', Id='rIdFootnotes', Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes', Target='footnotes.xml')
        parts['word/_rels/document.xml.rels'] = ET.tostring(rels)
        ctypes = ET.fromstring(parts['[Content_Types].xml'])
        ET.SubElement(ctypes, '{http://schemas.openxmlformats.org/package/2006/content-types}Override', PartName='/word/footnotes.xml', ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
        parts['[Content_Types].xml'] = ET.tostring(ctypes)
        parts['word/footnotes.xml'] = ('<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>각주 원문을 보존합니다</w:t></w:r></w:p></w:footnote></w:footnotes>').encode()
        data = io.BytesIO()
        with zipfile.ZipFile(data, 'w') as z:
            for name, value in parts.items(): z.writestr(name, value)
        self.docs['before'] = data.getvalue(); self.manifest['sources']['before']['sha256'] = hashlib.sha256(data.getvalue()).hexdigest()
        self.write_package(); a = load_package(self.package); out = self.root / 'out.docx'
        warnings = generate(a, initial_rows(a), out)
        cell = Document(out).tables[0].cell(2, 2)
        self.assertEqual(len(cell.tables), 1)
        self.assertIn('각주 원문을 보존합니다', cell.text)
        self.assertTrue(any('각주' in w for w in warnings))


if __name__ == "__main__": unittest.main()
